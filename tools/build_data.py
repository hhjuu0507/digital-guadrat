#!/usr/bin/env python3
"""
방형구 결과 엑셀 → 정적 열람 사이트용 파일 만들기

  python3 tools/build_data.py  <원본.xlsx>  <site 폴더>

만들어지는 파일 (site 폴더 안):
  data.js              사이트가 읽는 데이터 (이름·학번 가림 처리됨)
  results-masked.xlsx  '엑셀 다운로드' 버튼용 파일 (이름·학번 가림 처리됨)

※ 원본 엑셀에는 학생 실명이 들어 있으니 site 폴더(인터넷에 올리는 폴더) 안에 넣지 마세요.
필요: Python 3, openpyxl  (pip install openpyxl)
"""
import collections
import datetime
import json
import pathlib
import re
import sys
import zipfile

import openpyxl

# 엑셀의 '제출 일시'가 UTC로 저장돼 있어서 한국 시간(+9시간)으로 바꿉니다.
# 엑셀 시간이 이미 한국 시간이라면 0 으로 바꾸세요.
OFFSET = datetime.timedelta(hours=9)

PAIR = re.compile(r"(\d{4,6})\s*([가-힣]{2,5})")  # '10234홍길동', '10234 홍길동'
LONE_ID = re.compile(r"(?<![\d*])\d{4,6}(?![\d*])")  # 이름 없이 학번만 있는 경우


def mask_name(name):
    """홍길동 → 홍○동 / 이서 → 이○ / 남궁민수 → 남○○수"""
    return name[0] + "○" if len(name) < 3 else name[0] + "○" * (len(name) - 2) + name[-1]


def mask_id(num):
    """10234 → 102** (학년·반만 남김)"""
    return num[:3] + "*" * (len(num) - 3)


def mask_text(text):
    text = PAIR.sub(lambda m: f"{mask_id(m.group(1))} {mask_name(m.group(2))}", str(text))
    text = LONE_ID.sub(lambda m: mask_id(m.group(0)), text)
    # 가린 이름은 '이○범'처럼 한글이 한 글자씩 떨어져 있어요. 연속 2글자 이상이면 못 가린 것.
    if re.search(r"[가-힣]{2,}", text):
        raise ValueError(f"가리지 못한 이름이 남았어요 (형식을 확인하세요): {text!r}")
    return text


def main(src, out_dir):
    out = pathlib.Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    need = ["학번 이름", "조", "Site", "상태", "식물종(한글명)", "학명", "AI 참고용 초록색 비율(%)", "식별 신뢰도(%)", "종 상태", "과(Family)", "속(Genus)", "제출 일시"]
    missing = [n for n in need if n not in col]
    if missing:
        raise SystemExit(f"엑셀에 없는 열: {missing}")

    raw_rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    raw_rows = [r for r in raw_rows if any(v not in (None, "") for v in r)]

    # 제출 하나 = (학번 이름, 조, Site, 제출 일시)
    subs = collections.OrderedDict()
    for r in raw_rows:
        key = (r[col["학번 이름"]], r[col["조"]], r[col["Site"]], r[col["제출 일시"]])
        subs.setdefault(key, []).append(r)

    result = []
    for (name, group, site, when), rows in subs.items():
        kst = when + OFFSET
        result.append(
            {
                "name": mask_text(name),
                "group": int(re.sub(r"\D", "", group)),
                "site": int(re.sub(r"\D", "", site)),
                "status": rows[0][col["상태"]],
                "time": kst.strftime("%Y-%m-%dT%H:%M:%S"),
                "timeText": kst.strftime("%Y-%m-%d %H:%M"),
                "species": [
                    {
                        "ko": r[col["식물종(한글명)"]],
                        "sci": r[col["학명"]],
                        "family": r[col["과(Family)"]],
                        "genus": r[col["속(Genus)"]],
                        "ratio": round(float(r[col["AI 참고용 초록색 비율(%)"]]), 2),
                        "conf": int(r[col["식별 신뢰도(%)"]]),
                        "status": r[col["종 상태"]],
                    }
                    for r in rows
                ],
            }
        )
    result.sort(key=lambda s: s["time"])
    for i, s in enumerate(result, 1):
        s["id"] = f"s{i:03d}"
    result.sort(key=lambda s: s["time"], reverse=True)  # 최신순

    data = {
        "meta": {"submissions": len(result), "rows": len(raw_rows)},
        "submissions": result,
    }
    # fetch 없이도(파일을 더블클릭해서 열어도) 동작하도록 .js 파일로 저장
    data_text = "window.QUADRAT_DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    (out / "data.js").write_text(data_text, encoding="utf-8")

    # 이름 가린 엑셀
    c_name, c_time = col["학번 이름"], col["제출 일시"]
    for row in ws.iter_rows(min_row=2):
        if row[c_name].value not in (None, ""):
            row[c_name].value = mask_text(row[c_name].value)
        if isinstance(row[c_time].value, datetime.datetime):
            row[c_time].value = row[c_time].value + OFFSET
            row[c_time].number_format = "yyyy-mm-dd hh:mm:ss"
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.title = "방형구 식물 조사 결과 (이름 가림)"
    xlsx_path = out / "results-masked.xlsx"
    wb.save(xlsx_path)

    # ---- 가림 검사: 원본 이름/학번이 결과물에 그대로 남았는지 확인 ----
    secrets = set()
    for r in raw_rows:
        text = str(r[c_name])
        for num, nm in PAIR.findall(text):
            secrets.update([num, nm])
        secrets.update(re.findall(r"\d{5,6}", text))
    with zipfile.ZipFile(xlsx_path) as z:
        xlsx_text = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
    leaks = sorted(s for s in secrets if s in data_text or s in xlsx_text)
    if leaks:
        raise SystemExit(f"가림 실패! 결과물에 원본 값이 남아 있어요: {leaks}")

    print(f"제출 {len(result)}건 / 행 {len(raw_rows)}개")
    print(f"가림 검사 통과: 원본 이름·학번 {len(secrets)}개가 결과물에 남아 있지 않아요")
    print(f"만든 파일: {out / 'data.js'}, {xlsx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
